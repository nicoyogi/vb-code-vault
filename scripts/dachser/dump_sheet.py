#!/usr/bin/env python3
"""Dump one worksheet as the JSON the Anmerkung audit tool consumes.

SheetJS hands the app a worksheet keyed by A1 address, each cell `{v: value}`,
plus `ws['!ref']` covering the used range. openpyxl can hand Node the same
thing, so `scripts/dachser/audit_workbook.mjs` can run the real `resolveDachser`
and `processDachser` against a workbook without shipping SheetJS to Node.

    py scripts/dachser/dump_sheet.py "D:\\TESTER\\Dachser\\<file>.xlsx" [sheet]

Formulas are read for their cached value (`data_only=True`), which is what
SheetJS reports. A cell that Excel would show blank is omitted entirely.
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def main():
    if len(sys.argv) < 2:
        print("usage: dump_sheet.py <workbook.xlsx> [sheet]", file=sys.stderr)
        return 2
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    path = Path(sys.argv[1])
    if not path.exists():
        print(f"error: no such workbook: {path}", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = sys.argv[2] if len(sys.argv) > 2 else wb.active.title
    if name not in wb.sheetnames:
        print(f"error: no sheet {name!r} in {wb.sheetnames}", file=sys.stderr)
        return 2

    ws = wb[name]
    cells = {}
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None or v == "":
                continue
            cells[cell.coordinate] = v

    if ws.max_row and ws.max_column:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    else:
        ref = "A1:A1"

    json.dump({"path": str(path), "sheet": name, "ref": ref, "cells": cells},
              sys.stdout, ensure_ascii=False, default=str)
    return 0


if __name__ == "__main__":
    sys.exit(main())
