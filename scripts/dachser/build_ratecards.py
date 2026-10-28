#!/usr/bin/env python3
"""Build the local-only Dachser ratecard asset from the two xlsx sources.

Output: assets/dachser-ratecards.js  (gitignored business data).
The shipped, public artifact is the encrypted bundle
assets/dachser-ratecards.enc.json, produced from that plaintext by
scripts/encrypt-dachser-ratecards.mjs.

    python scripts/dachser/build_ratecards.py

Two matrices are emitted, because the Sonderfahrt / bisherigen rule compares
against both:

  NEW  Preismatrix_Dornstadt_International ab 01.07.2026.xlsx
       sheet "Preismatrix Dornstadt", headers row 1, data from row 3.
  OLD  DACHSER_Int. 01.07.24-30.06.26.xlsx
       sheet "Preismatrix Dornst.,Geisl.,Eisl", headers row 11, data from row 13.

Layout (both): column A = "Von" kg, B = "Bis" kg, C = "Staffel" label, and
columns D.. carry one lane each, headed by the country label from the header
row. Each matrix is a list of weight brackets, each bracket a lane->price map.

A third key, `national`, carries the domestic card (nine DE zones by weight
bracket) plus the two-digit PLZ-prefix -> zone table, which the Referenz count
needs. It is not a matrix: its columns are zones, not lanes.
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "assets" / "dachser-ratecards.js"

SHEETS = {
    "new": {
        "path": r"D:\TESTER\Preismatrix_Dornstadt_International ab 01.07.2026.xlsx",
        "sheet": "Preismatrix Dornstadt",
        "header_row": 1,
        "first_data_row": 3,
        "source": "Preismatrix_Dornstadt_International ab 01.07.2026.xlsx",
    },
    "old": {
        "path": r"D:\TESTER\DACHSER_Int. 01.07.24-30.06.26.xlsx",
        "sheet": "Preismatrix Dornst.,Geisl.,Eisl",
        "header_row": 11,
        "first_data_row": 13,
        "source": "DACHSER_Int. 01.07.24-30.06.26.xlsx",
    },
}


NATIONAL = {
    "path": r"D:\TESTER\Dachser national ratecard.xlsx",
    "sheet": "Basis",
    "header_row": 1,
    "first_data_row": 2,
    "source": "Dachser national ratecard.xlsx",
}
ZONES = {
    "path": r"D:\TESTER\Zoneneinteilung_national_international_01.07.24.xlsx",
    "sheet": "DE",
    "source": "Zoneneinteilung_national_international_01.07.24.xlsx",
}
# The supplied workbook loses its decimal separator above 2400 kg (DE1 reads
# 220759 at 2600 kg after 210.91 at 2400). A cell whose step over this factor
# gets its separator back; a bracket that still overflows is withheld.
MAX_STEP_FACTOR = 1.5


def build(spec):
    wb = openpyxl.load_workbook(spec["path"], data_only=True)
    ws = wb[spec["sheet"]]
    heads = {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=spec["header_row"], column=c).value
        if h is not None and str(h).strip():
            heads[c] = str(h).strip()

    brackets = []
    for r in range(spec["first_data_row"], ws.max_row + 1):
        von = ws.cell(row=r, column=1).value
        bis = ws.cell(row=r, column=2).value
        if von is None or bis is None:
            continue
        try:
            von_i, bis_i = int(von), int(bis)
        except (TypeError, ValueError):
            continue
        lanes = {}
        for c, label in heads.items():
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)):
                lanes[label] = round(float(v), 2)
        brackets.append({"von": von_i, "bis": bis_i, "lanes": lanes})
    return {
        "schema": "dachser-ratecard/v1",
        "source": spec["source"],
        "brackets": brackets,
    }


def build_national(spec):
    """Nine zone columns, one row per weight bracket. Above 2400 kg the supplied
    workbook drops the decimal separator (DE1 reads 220759 at 2600 kg), so a cell
    whose step overflows MAX_STEP_FACTOR is retried at value/1000, but only when
    that still reaches the previous bracket's cell; a bracket is withheld only if
    it overflows after that repair."""
    wb = openpyxl.load_workbook(spec["path"], data_only=True)
    ws = wb[spec["sheet"]]
    zones_at = {}
    for c in range(2, ws.max_column + 1):
        label = ws.cell(row=spec["header_row"], column=c).value
        if label is not None and str(label).strip():
            zones_at[c] = str(label).strip()

    brackets = []
    capped_at_bis = None
    recovered_from_bis = None
    withheld_from_bis = None
    for r in range(spec["first_data_row"], ws.max_row + 1):
        bis = ws.cell(row=r, column=1).value
        try:
            bis_i = int(bis)
        except (TypeError, ValueError):
            continue
        prev = brackets[-1]["zones"] if brackets else None
        zones = {}
        repaired = False
        for c, label in zones_at.items():
            v = ws.cell(row=r, column=c).value
            if not isinstance(v, (int, float)):
                continue
            v = float(v)
            if prev and prev.get(label):
                step = v / prev[label]
                divided = v / 1000.0
                if (step > MAX_STEP_FACTOR
                        and round(divided, 2) >= prev[label]
                        and divided / prev[label] <= MAX_STEP_FACTOR):
                    v = divided
                    repaired = True
            zones[label] = round(v, 2)
        if not zones:
            continue
        if prev:
            step = max(zones[k] / prev[k] for k in zones if k in prev and prev[k])
            if step > MAX_STEP_FACTOR:
                if capped_at_bis is None:
                    capped_at_bis = brackets[-1]["bis"]
                withheld_from_bis = bis_i
                break
        if repaired and recovered_from_bis is None:
            capped_at_bis = brackets[-1]["bis"]
            recovered_from_bis = bis_i
        brackets.append({"bis": bis_i, "zones": zones})
    if capped_at_bis is None:
        capped_at_bis = brackets[-1]["bis"]
    return {
        "schema": "dachser-national-ratecard/v1",
        "source": spec["source"],
        "plz_source": ZONES["source"],
        "capped_at_bis": capped_at_bis,
        "recovered_from_bis": recovered_from_bis,
        "withheld_from_bis": withheld_from_bis,
        "brackets": brackets,
    }


def build_plz(spec):
    """Two-digit Leitcode prefix -> zone digit, from the DE sheet (B = zone,
    C = Leitcode). Zero padding matters: the engine slices the PLZ to two chars."""
    wb = openpyxl.load_workbook(spec["path"], data_only=True)
    ws = wb[spec["sheet"]]
    plz = {}
    for r in range(2, ws.max_row + 1):
        zone = ws.cell(row=r, column=2).value
        leit = ws.cell(row=r, column=3).value
        if zone is None or leit is None:
            continue
        digits = "".join(ch for ch in str(zone) if ch.isdigit())
        if not digits:
            continue
        try:
            plz["%02d" % int(leit)] = digits
        except (TypeError, ValueError):
            continue
    return dict(sorted(plz.items()))


def build_plz_missing(plz):
    """The two-digit prefixes the zone sheet does not carry at all, so a missing
    one is visible instead of silent."""
    return ["%02d" % i for i in range(100) if "%02d" % i not in plz]


def emit_js(new_card, old_card, national):
    lines = []
    lines.append("/* Dachser ratecards - generated from")
    lines.append("     %s" % new_card["source"])
    lines.append("     %s" % old_card["source"])
    lines.append("   by scripts/dachser/build_ratecards.py. Business data: local only,")
    lines.append("   gitignored; the public artifact is assets/dachser-ratecards.enc.json. */")
    lines.append("globalThis.DACHSER_RATECARDS = {")
    lines.append('  schema: "dachser-ratecards/v1",')
    for key, card in (("new", new_card), ("old", old_card)):
        lines.append("  %s: {" % key)
        lines.append("    source: %s," % json.dumps(card["source"]))
        lines.append("    brackets: [")
        for b in card["brackets"]:
            lines.append("      { von: %d, bis: %d, lanes: %s }," % (
                b["von"], b["bis"], json.dumps(b["lanes"], ensure_ascii=False)))
        lines.append("    ],")
        lines.append("  },")
    lines.append("  national: {")
    lines.append("    schema: %s," % json.dumps(national["schema"]))
    lines.append("    source: %s," % json.dumps(national["source"]))
    lines.append("    plz_source: %s," % json.dumps(national["plz_source"]))
    lines.append("    capped_at_bis: %s," % json.dumps(national["capped_at_bis"]))
    lines.append("    recovered_from_bis: %s," % json.dumps(national["recovered_from_bis"]))
    lines.append("    withheld_from_bis: %s," % json.dumps(national["withheld_from_bis"]))
    lines.append("    brackets: [")
    for b in national["brackets"]:
        lines.append("      { bis: %d, zones: %s }," % (
            b["bis"], json.dumps(b["zones"], ensure_ascii=False)))
    lines.append("    ],")
    lines.append("    plz: %s," % json.dumps(national["plz"], ensure_ascii=False))
    lines.append("    plz_missing: %s," % json.dumps(national["plz_missing"]))
    lines.append("  },")
    lines.append("};")
    lines.append("")
    return "\n".join(lines)


def main():
    cards = {}
    for key, spec in SHEETS.items():
        p = Path(spec["path"])
        if not p.exists():
            sys.exit("source workbook not found: %s" % p)
        cards[key] = build(spec)
        print("%-4s %2d brackets, %d lanes in first bracket" % (
            key, len(cards[key]["brackets"]),
            len(cards[key]["brackets"][0]["lanes"]) if cards[key]["brackets"] else 0))

    for spec in (NATIONAL, ZONES):
        if not Path(spec["path"]).exists():
            sys.exit("source workbook not found: %s" % spec["path"])
    national = build_national(NATIONAL)
    national["plz"] = build_plz(ZONES)
    national["plz_missing"] = build_plz_missing(national["plz"])
    print("nat  %2d brackets (sound to %d kg, recovered from %s kg, withheld from %s), %d PLZ prefixes, missing %s" % (
        len(national["brackets"]), national["capped_at_bis"],
        national["recovered_from_bis"], national["withheld_from_bis"], len(national["plz"]),
        ",".join(national["plz_missing"]) or "none"))
    OUT.write_text(emit_js(cards["new"], cards["old"], national), encoding="utf-8")
    print("Wrote %s (%d bytes)" % (OUT, OUT.stat().st_size))


if __name__ == "__main__":
    main()
