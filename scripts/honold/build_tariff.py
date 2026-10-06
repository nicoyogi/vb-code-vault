#!/usr/bin/env python3
"""Build the local-only Honold OLD tariff asset from its xlsx source.

Output: assets/honold-tariff.js  (gitignored business data).
The shipped, public artifact is the encrypted bundle
assets/honold-tariff.enc.json, produced from that plaintext by
scripts/encrypt-honold-tariff.mjs.

    python scripts/honold/build_tariff.py [path/to/"Honold OLD Tariff.xlsx"]

Sheet1 layout: row 1 = zone headers (B..BE), column A from row 2 = "Bis"
weight bracket, cells = freight price in EUR.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_XLSX = r"D:\TESTER\Honold OLD Tariff.xlsx"
OUT = ROOT / "assets" / "honold-tariff.js"

# Headers that dropped their leading country code in the source sheet.
CODE_ALIASES = {"Spanien 3": "ES"}


def country_code(label):
    m = re.match(r"^([A-Z]{2})", label)
    if m:
        return m.group(1)
    return CODE_ALIASES.get(label)


def build(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Sheet1"]
    tiers = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None or v == "":
            break
        tiers.append(v)
    zones = {}
    codes = {}
    for c in range(2, ws.max_column + 1):
        label = ws.cell(1, c).value
        if label is None or label == "":
            continue
        label = str(label).strip()
        rates = []
        for i in range(len(tiers)):
            rates.append(ws.cell(2 + i, c).value)
        if all(v is None for v in rates):
            continue
        zones[label] = rates
        code = country_code(label)
        if code:
            codes.setdefault(code, []).append(label)
    return {"tiers": tiers, "zones": zones, "codes": codes}


def emit_js(data, source_name):
    lines = []
    lines.append("/* Honold OLD tariff - generated from %s (Sheet1) by" % source_name)
    lines.append("   scripts/honold/build_tariff.py. Business data: local only, gitignored;")
    lines.append("   the public artifact is assets/honold-tariff.enc.json. */")
    lines.append("globalThis.HONOLD_TARIFF = {")
    lines.append('  schema: "honold-old-tariff/v1",')
    lines.append('  source: %s,' % json.dumps(source_name))
    lines.append("  tiers: %s," % json.dumps(data["tiers"]))
    lines.append("  zones: {")
    items = list(data["zones"].items())
    for i, (label, rates) in enumerate(items):
        comma = "," if i < len(items) - 1 else ""
        lines.append("    %s: %s%s" % (json.dumps(label, ensure_ascii=False), json.dumps(rates), comma))
    lines.append("  },")
    lines.append("  codes: %s" % json.dumps(data["codes"], ensure_ascii=False))
    lines.append("};")
    lines.append("")
    return "\n".join(lines)


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    xlsx_path = Path(xlsx)
    if not xlsx_path.exists():
        sys.exit("source workbook not found: %s" % xlsx_path)
    data = build(xlsx_path)
    OUT.write_text(emit_js(data, xlsx_path.name), encoding="utf-8")
    print("wrote %s: %d tiers x %d zones, codes=%s"
          % (OUT, len(data["tiers"]), len(data["zones"]), sorted(data["codes"])))


if __name__ == "__main__":
    main()
